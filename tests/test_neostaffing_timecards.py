"""Canonical timecards, leadership scope, archive acknowledgement and retention."""
import io
import json
import unittest
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from unittest.mock import patch

from sqlalchemy import event

from app.extensions import db
from app.models import StaffingDailyAttendance
from app.models.staffing_timecard import StaffingTimecardSlice as Slice, StaffingTimecardSegment as Segment, StaffingTimecardWeek as Week, StaffingTimecardEdit as Edit
from app.services import neostaffing as staffing
from app.services import neostaffing_timecards as tc
from app.services import neostaffing_timecard_exports as exports
from tests import test_neosektor_employees as fixtures


class TimecardsTest(unittest.TestCase):
    setUp = fixtures.SektorEmployeesTest.setUp
    setUpBase = fixtures.SektorEmployeesTest.setUpBase
    tearDown = fixtures.SektorEmployeesTest.tearDown
    _login_approved_user = fixtures.SektorEmployeesTest._login_approved_user
    _add_sort_operation = fixtures.SektorEmployeesTest._add_sort_operation
    person = fixtures.SektorEmployeesTest.person
    active_sort = fixtures.SektorEmployeesTest.active_sort
    page = fixtures.SektorEmployeesTest.page
    form = fixtures.SektorEmployeesTest.form

    def attendance(self, person=None, status="here"):
        person = person or self.workers["ebm"]
        if not staffing.current_night_attendance_operation():
            self.active_sort()
        staffing.save_operational_manage_attendance(self.form(person, status), self.user,
            [self.areas["ebm"].id], form_submission=True, home_only=True)
        db.session.commit()
        return Slice.query.filter_by(person_id=person.id).one()

    def command(self, row, segments=None):
        return {"id": row.id, "version": row.version,
            "segments": segments if segments is not None else [{"start":"2026-09-12T22:00:00-05:00", "end":"2026-09-13T02:10:00-05:00"}]}

    def test_workday_week_and_exact_multiple_segments(self):
        row = self.attendance()
        self.assertEqual(row.workday_date, date(2026, 9, 12))
        self.assertEqual(tc.week_start(row.workday_date), date(2026, 9, 6))
        self.assertEqual(tc.week_start(date(2026, 9, 13)), date(2026, 9, 13))
        tc.save_segments(self.user, [self.command(row)], as_of=row.workday_date)
        db.session.commit()
        result = tc.read_rows(self.user, row.workday_date, row.workday_date)[0]
        self.assertEqual(result["hours"], Decimal("4.17"))
        parts = [{"start":"2026-09-12T22:00-05:00", "end":"2026-09-12T23:00-05:00"},
                 {"start":"2026-09-13T01:00-05:00", "end":"2026-09-13T02:10-05:00"}]
        tc.save_segments(self.user, [self.command(row, parts)], as_of=row.workday_date)
        db.session.commit()
        self.assertEqual(tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["hours"], Decimal("2.17"))
        self.assertEqual(Segment.query.count(), 2)
        self.assertEqual(Edit.query.count(), 3)

    def test_attendance_gating_retains_audit_and_stale_rejection(self):
        row = self.attendance()
        stale = self.command(row)
        tc.save_segments(self.user, [stale], as_of=row.workday_date)
        db.session.commit()
        with self.assertRaisesRegex(ValueError, "changed"):
            tc.save_segments(self.user, [stale], as_of=row.workday_date)
        db.session.rollback()
        self.attendance(status="call_in")
        self.assertIsNone(tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["hours"])
        self.assertEqual(Segment.query.count(), 1)
        self.assertEqual(Edit.query.count(), 3)
        with self.assertRaisesRegex(ValueError, "Here"):
            tc.save_segments(self.user, [self.command(row)], as_of=row.workday_date)
        db.session.rollback()
        self.attendance(status="here")
        self.assertEqual(tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["hours"], Decimal("4.17"))

    def test_role_alone_never_edits_and_different_employees_independent(self):
        row, peer = self.attendance(), self.attendance(self.peer)
        a, b = self.command(row), self.command(peer)
        tc.save_segments(self.user, [a], as_of=row.workday_date)
        db.session.commit()
        tc.save_segments(self.user, [b], as_of=row.workday_date)
        db.session.commit()
        self.manager.active = False
        db.session.commit()
        with self.assertRaisesRegex(ValueError, "management"):
            tc.save_segments(self.user, [self.command(row)], as_of=row.workday_date)
        db.session.rollback()
        self.assertEqual(Segment.query.count(), 2)

    def test_exceptions_and_no_get_writes(self):
        row = self.attendance()
        self.assertIn("Here without times", tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["issues"])
        parts = [{"start":"2026-09-12T22:00-05:00", "end":"2026-09-13T02:00-05:00"},
                 {"start":"2026-09-13T01:00-05:00", "end":"2026-09-13T03:00-05:00"}, {"start":"2026-09-13T04:00-05:00"}]
        tc.save_segments(self.user, [self.command(row, parts)], as_of=row.workday_date)
        db.session.commit()
        day = row.workday_date
        sql = []
        def capture(_c, _cur, statement, *_args): sql.append(statement.lstrip().lower())
        event.listen(db.engine, "before_cursor_execute", capture)
        try:
            result = tc.read_rows(self.user, day, day)[0]
        finally:
            event.remove(db.engine, "before_cursor_execute", capture)
        self.assertIn("Overlapping segments", result["issues"])
        self.assertIn("Incomplete segment", result["issues"])
        self.assertFalse(any(statement.startswith(("insert", "update", "delete")) for statement in sql))
        # Includes one set-based Combo-config lookup for missing-half exceptions.
        self.assertLessEqual(sum(statement.startswith("select") for statement in sql), 7)

    def test_reports_sql_grouping_and_node_mode(self):
        row, peer = self.attendance(), self.attendance(self.peer)
        tc.save_segments(self.user, [self.command(row), self.command(peer)], as_of=row.workday_date)
        db.session.commit()
        report = tc.report_context(self.user, {"date": str(row.workday_date), "view":"sort"}, as_of=row.workday_date)
        self.assertEqual(report["total"], 1)
        self.assertEqual(report["report"][0][5], Decimal("8.33"))
        report = tc.report_context(self.user, {"date": str(row.workday_date), "view":"employee"}, as_of=row.workday_date)
        self.assertEqual(report["total"], 2)
        response = self.client.get('/neosektor/manage-employees?area=ebm&view=all&mode=times')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'SAVE CHANGED TIMES', response.data)
        self.assertIn(self.peer.full_name.encode(), response.data)
        self.assertEqual(self.client.get('/neostaffing/timecards?date=2026-09-12').status_code, 200)

    def test_archive_scope_acknowledgement_and_invalidation(self):
        row = self.attendance()
        tc.save_segments(self.user, [self.command(row)], as_of=row.workday_date)
        db.session.commit()
        week = tc.week_start(row.workday_date)
        package, token = exports.generate(self.user, week, complete=False, as_of=date(2026,9,14))
        self.assertIsNone(token)
        self.assertIsNone(db.session.get(Week, week).purge_after)
        with self.assertRaisesRegex(ValueError, "Master"):
            exports.generate(self.user, week, complete=True, as_of=date(2026,9,14))
        with patch.object(exports, "can_archive", return_value=True):
            package, token = exports.generate(self.user, week, complete=True, as_of=date(2026,9,14))
            self.assertIsNone(db.session.get(Week, week).purge_after)
            exports.acknowledge_download(self.user, token, now=datetime(2026,9,14))
            db.session.commit()
            self.assertEqual(db.session.get(Week, week).purge_after, datetime(2026,9,17))
            exports.acknowledge_download(self.user, token, now=datetime(2026,9,15))
            self.assertEqual(db.session.get(Week, week).purge_after, datetime(2026,9,17))
            tc.save_segments(self.user, [self.command(row, [])], as_of=date(2026,9,15))
            db.session.commit()
            self.assertIsNone(db.session.get(Week, week).purge_after)
            with self.assertRaisesRegex(ValueError, "stale"):
                exports.acknowledge_download(self.user, token)
            db.session.rollback()
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            self.assertEqual(set(archive.namelist()), {"timecards.xlsx", "timecards.csv", "edit-history.csv", "manifest.json"})
            manifest = json.loads(archive.read("manifest.json"))
            self.assertTrue(manifest["complete"])
            self.assertEqual(manifest["slice_count"], 1)
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(archive.read("timecards.xlsx")))
            self.assertEqual(book.sheetnames, ["Summary", "All Employee Hours", "Employee Weekly", "Sort Labor", "Exceptions"])
            self.assertEqual(book["All Employee Hours"].freeze_panes, "C5")
            self.assertEqual(book["All Employee Hours"]["K5"].value, 4.17)
            self.assertEqual(len(book["All Employee Hours"].tables), 1)

    def test_bounded_purge_and_one_calendar_month(self):
        row, peer = self.attendance(), self.attendance(self.peer)
        self.assertEqual(tc.retention_cutoff(date(2026,3,31)), date(2026,2,28))
        with patch.object(exports, "can_archive", return_value=True):
            _, token = exports.generate(self.user, date(2026,9,6), complete=True, as_of=date(2026,9,14))
            exports.acknowledge_download(self.user, token, now=datetime(2026,9,14))
            db.session.commit()
        self.assertEqual(tc.cleanup(as_of=date(2026,9,16), now=datetime(2026,9,16)), 0)
        self.assertEqual(tc.cleanup(as_of=date(2026,9,17), now=datetime(2026,9,17), batch_size=1), 1)
        db.session.commit()
        self.assertEqual(Slice.query.count(), 1)
        self.assertEqual(tc.cleanup(as_of=date(2026,9,17), now=datetime(2026,9,17), batch_size=1), 1)
        db.session.commit()
        self.assertEqual(Slice.query.count(), 0)
        self.assertEqual(Edit.query.count(), 0)
        self.assertEqual(StaffingDailyAttendance.query.count(), 2)
        self.assertEqual(db.session.get(Week, date(2026,9,6)).purged_slices, 2)

    def test_combo_cross_sort_authority_labor_and_employee_aggregation(self):
        from app.models import StaffingUnit, StaffingWorkAssignment
        from app.models.staffing_accountability import StaffingComboWorkday
        worker = self.workers["ebm"]
        worker.classification = "full_time_combo"
        other_sort = StaffingUnit(name="Configured First", unit_type="sort")
        other_op = StaffingUnit(name="Other Operation", unit_type="operation", parent=other_sort)
        other_dept = StaffingUnit(name="Other Department", unit_type="department", parent=other_op)
        other_area = StaffingUnit(name="Other Area", unit_type="work_area", parent=other_dept)
        db.session.add_all([other_sort, other_op, other_dept, other_area])
        db.session.flush()
        staffing.assign_work_area(worker, other_area)
        db.session.add(StaffingComboWorkday(person_id=worker.id, first_sort_id=other_sort.id, second_sort_id=self.night.id))
        db.session.commit()
        row = self.attendance()
        operation = self._add_sort_operation(row.workday_date, "configured first")
        assignment = StaffingWorkAssignment.query.filter_by(person_id=worker.id, work_area_unit_id=other_area.id, active=True).one()
        tc.sync_attendance(operation, {worker.id:"here"}, {worker.id:assignment}, self.user.id, as_of=row.workday_date)
        db.session.commit()
        other = Slice.query.filter_by(sort_unit_id=other_sort.id).one()
        tc.save_segments(self.user, [self.command(row), self.command(other, [{"start":"2026-09-12T15:00-05:00", "end":"2026-09-12T17:00-05:00"}])], as_of=row.workday_date)
        db.session.commit()
        result = tc.report_context(self.user, {"date":str(row.workday_date), "view":"employee"}, as_of=row.workday_date)
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["report"][0][2], Decimal("6.17"))
        result = tc.report_context(self.user, {"date":str(row.workday_date), "view":"sort"}, as_of=row.workday_date)
        self.assertEqual(result["total"], 2)
        snapshot = json.loads(other.context_json)
        snapshot["employee"] = "Renamed employee"
        other.context_json = json.dumps(snapshot)
        db.session.commit()
        with patch.object(exports, "can_archive", return_value=True):
            package, _ = exports.generate(self.user, tc.week_start(row.workday_date), complete=True, as_of=date(2026,9,14))
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            self.assertEqual(json.loads(archive.read("manifest.json"))["employee_count"], 1)
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(archive.read("timecards.xlsx")))
            self.assertEqual(book["Employee Weekly"]["C5"].value, 6.17)
        # Removing configured pairing removes the cross-Sort delegation.
        db.session.delete(db.session.get(StaffingComboWorkday, worker.id))
        db.session.commit()
        with self.assertRaisesRegex(ValueError, "scope"):
            tc.save_segments(self.user, [self.command(other, [])], as_of=row.workday_date)
        db.session.rollback()

    def test_bootstrap_idempotent_and_retention_independent_of_archive(self):
        row = self.attendance()
        Edit.query.delete()
        Slice.query.delete()
        db.session.commit()
        self.assertEqual(tc.backfill_retained_attendance(date(2026,9,12)), 1)
        db.session.commit()
        self.assertEqual(tc.backfill_retained_attendance(date(2026,9,12)), 0)
        self.assertEqual(Slice.query.one().attendance_status, "here")
        self.assertEqual(tc.cleanup(as_of=date(2026,10,13)), 1)
        db.session.commit()
        self.assertEqual(tc.backfill_retained_attendance(date(2026,10,13)), 0)
        self.assertEqual(StaffingDailyAttendance.query.count(), 1)

    def test_clock_shortcuts_and_dst_validation(self):
        row = self.attendance()
        tc.save_segments(self.user, [self.command(row, [{"start":"22:00", "end":"02:10"}])], as_of=row.workday_date)
        db.session.commit()
        self.assertEqual(tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["hours"], Decimal("4.17"))
        with self.assertRaises(ValueError):
            tc.parse_timestamp("2026-11-01T01:30", "America/Chicago")
        with self.assertRaises(ValueError):
            tc.parse_timestamp("2026-03-08T02:30", "America/Chicago")
        self.assertNotEqual(tc.parse_timestamp("2026-11-01T01:30-05:00", "America/Chicago"),
                            tc.parse_timestamp("2026-11-01T01:30-06:00", "America/Chicago"))

    def test_complete_export_includes_outside_scope_and_nonhere_status(self):
        row = self.attendance()
        from app.models import StaffingWorkAssignment
        assignment = StaffingWorkAssignment.query.filter_by(person_id=self.outsider.id, active=True).one()
        tc.sync_attendance(self.operation, {self.outsider.id:"call_in"}, {self.outsider.id:assignment}, self.user.id, as_of=row.workday_date)
        db.session.commit()
        self.assertEqual(len(tc.read_rows(self.user, row.workday_date, row.workday_date)), 1)
        with patch.object(exports, "can_archive", return_value=True):
            package, _ = exports.generate(self.user, tc.week_start(row.workday_date), complete=True, as_of=date(2026,9,14))
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            self.assertEqual(json.loads(archive.read("manifest.json"))["slice_count"], 2)
            self.assertIn(b"call_in", archive.read("timecards.csv"))
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(archive.read("timecards.xlsx")))
            self.assertIsNone(book["All Employee Hours"]["K6"].value)
        receipt = tc.lock_week(row.workday_date)
        receipt.purge_after = datetime(2026,9,17)
        db.session.commit()
        # A status-only attendance correction must invalidate the whole week.
        self.attendance(status="call_in")
        self.assertIsNone(tc.lock_week(row.workday_date).purge_after)

    def test_empty_archived_weeks_cannot_starve_cleanup(self):
        for offset in range(9):
            receipt = tc.lock_week(date(2026, 9, 6) - timedelta(weeks=offset))
            receipt.purge_after = datetime(2026, 9, 17)
        db.session.commit()
        self.assertEqual(tc.cleanup(as_of=date(2026,9,17), now=datetime(2026,9,17)), 0)
        db.session.commit()
        self.assertEqual(Week.query.filter(Week.purge_after.is_not(None)).count(), 1)
        tc.cleanup(as_of=date(2026,9,17), now=datetime(2026,9,17))
        db.session.commit()
        self.assertEqual(Week.query.filter(Week.purge_after.is_not(None)).count(), 0)

    def test_noop_preserves_version_and_purge_and_subsecond_aggregation(self):
        row = self.attendance()
        command = self.command(row, [{"start":"2026-09-12T22:00:00.100000-05:00",
                                     "end":"2026-09-12T22:00:18.099999-05:00"}])
        tc.save_segments(self.user, [command], as_of=row.workday_date)
        db.session.commit()
        self.assertEqual(tc.read_rows(self.user, row.workday_date, row.workday_date)[0]["hours"], Decimal("0.00"))
        report = tc.report_context(self.user, {"date":str(row.workday_date), "view":"employee"}, as_of=row.workday_date)
        self.assertEqual(report["report"][0][2], Decimal("0.00"))
        receipt = tc.lock_week(row.workday_date)
        receipt.purge_after = datetime(2026,9,17)
        db.session.commit()
        version = row.version
        tc.save_segments(self.user, [self.command(row, command["segments"])], as_of=row.workday_date)
        db.session.commit()
        self.assertEqual(row.version, version)
        self.assertEqual(tc.lock_week(row.workday_date).purge_after, datetime(2026,9,17))
        parts = [{"start":"2026-09-12T22:00:00.999999-05:00", "end":"2026-09-12T22:00:19-05:00"}]
        tc.save_segments(self.user, [self.command(row, parts)], as_of=row.workday_date)
        db.session.commit()
        report = tc.report_context(self.user, {"date":str(row.workday_date), "view":"employee"}, as_of=row.workday_date)
        self.assertEqual(report["report"][0][2], Decimal("0.01"))

    def test_schema_creation_is_additive_and_idempotent(self):
        from sqlalchemy import inspect
        from app.services.schema_sync import _create_missing_application_tables
        person_id = self.peer.id
        db.session.commit()  # Release fixture reads before separate-connection DDL.
        for model in (Edit, Segment, Slice, Week):
            model.__table__.drop(db.engine)
        for _ in range(2):
            _create_missing_application_tables(set(inspect(db.engine).get_table_names()))
            db.session.commit()
        self.assertEqual(self.peer.id, person_id)
        self.assertEqual(StaffingDailyAttendance.query.count(), 0)
        self.assertEqual(self.attendance().attendance_status, "here")

    def test_clean_slices_do_not_bury_exceptions_and_overlap_survives_pagination(self):
        row, peer = self.attendance(), self.attendance(self.peer)
        tc.save_segments(self.user, [self.command(row)], as_of=row.workday_date)
        db.session.commit()
        report = tc.report_context(self.user, {"date":str(row.workday_date), "view":"exceptions"}, as_of=row.workday_date)
        self.assertEqual(report["total"], 1)
        self.assertEqual(report["report"][0][-1], "Here without times")
        # A second Sort outside the selected slice still participates in the
        # employee-wide overlap check, without requiring it on the same page.
        from app.models import StaffingUnit
        other_sort = StaffingUnit(name="Other slice", unit_type="sort")
        db.session.add(other_sort)
        db.session.flush()
        other = Slice(person_id=row.person_id, workday_date=row.workday_date,
            sort_unit_id=other_sort.id, sort_date_operation_id=row.sort_date_operation_id,
            work_area_unit_id=row.work_area_unit_id, context_json=row.context_json,
            attendance_status="here")
        db.session.add(other)
        db.session.flush()
        db.session.add(Segment(slice_id=other.id, start_utc=datetime(2026,9,13,4), end_utc=datetime(2026,9,13,6)))
        db.session.commit()
        result = tc.read_rows(self.user, row.workday_date, row.workday_date, slice_ids=[row.id])[0]
        self.assertIn("Overlapping segments", result["issues"])

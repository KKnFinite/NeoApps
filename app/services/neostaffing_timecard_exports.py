"""In-memory XLSX/CSV package; only an acknowledged complete export arms purge."""
import csv
import hashlib
import io
import json
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta

from flask import current_app
from itsdangerous import URLSafeTimedSerializer, BadSignature

from app.extensions import db
from app.models.staffing_timecard import StaffingTimecardEdit as Edit
from app.services.access_control import user_can_access_app
from app.services import neostaffing_timecards as timecards


def can_archive(user):
    return user_can_access_app(user, "neostaffing", minimum_role="master")


def serializer():
    return URLSafeTimedSerializer(current_app.config["SECRET_KEY"], salt="staffing-complete-week-v1")


def safe_text(value):
    # CSV/Excel formula injection protection for employee and organization text.
    return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value


def csv_bytes(headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows([[safe_text(value) for value in row] for row in rows])
    return output.getvalue().encode("utf-8-sig")


def build_package(rows, week, version, *, complete, generated_at=None, edits=(), previously_purged_slices=0):
    """Python production exporter, no desktop runtime or archive DB blobs."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    generated_at = generated_at or datetime.utcnow()
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = ["Employee ID", "Employee", "Workday", "Gateway", "Sort", "Operation", "Department", "Work Area",
               "Attendance", "Worked Minutes", "Hours", "Slice ID", "Source Operation ID", "Version", "Person ID"]
    detail, raw, exceptions = [], [], []
    weekly, labor = defaultdict(lambda: [0, set()]), defaultdict(lambda: [0, set()])
    employee_labels = {}
    for item in rows:
        row = item["slice"]
        prefix = [item["employee_id"], item["employee"], row.workday_date, item["gateway"], item["sort"],
                  item["operation"], item["department"], item["work_area"]]
        detail.append(prefix + [row.attendance_status or "Unmarked", item["seconds"] / 60 if item["seconds"] is not None else None,
                      float(item["hours"]) if item["hours"] is not None else None, row.id, row.sort_date_operation_id, row.version, row.person_id])
        for part in item["segments"] or [None]:
            raw.append(prefix + [row.attendance_status, row.id, row.sort_date_operation_id, row.version,
                part.id if part else None, part.start_utc.isoformat() + "Z" if part and part.start_utc else None,
                part.end_utc.isoformat() + "Z" if part and part.end_utc else None,
                item["timezone"], row.updated_by_user_id, row.updated_at.isoformat() + "Z", row.person_id])
        for issue in item["issues"]:
            exceptions.append(prefix + [issue, row.id])
        employee_key = row.person_id
        employee_labels[employee_key] = (item["employee_id"], item["employee"])
        labor_key = (item["gateway"], item["sort"], item["operation"], item["department"], item["work_area"])
        for target, key in ((weekly, employee_key), (labor, labor_key)):
            target[key][0] += item["seconds"] or 0
            target[key][1].add(row.attendance_status or "Unmarked")

    def sheet(name, titles, values, date_columns=(), decimal_columns=()):
        page = workbook.create_sheet(name)
        page.sheet_view.showGridLines = False
        page.append([name])
        page["A1"].font = Font(name="Arial", size=14, bold=True, color="182C3B")
        page.append([f"Week {week.isoformat()} through {(week + timedelta(days=6)).isoformat()}"])
        page.append([])
        page.append(titles)
        for record in values:
            page.append([safe_text(value) for value in record])
        for cells in page.iter_rows(min_row=4):
            for cell in cells:
                cell.font = Font(name="Arial", size=10, color="172536")
                cell.alignment = Alignment(vertical="center")
        for cell in page[4]:
            cell.fill = PatternFill("solid", fgColor="233A4A")
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, title in enumerate(titles, 1):
            width = max(len(str(title)) + 3, 15)
            if "Employee" == title or "Exception" == title:
                width = 30
            page.column_dimensions[get_column_letter(index)].width = min(width, 36)
        for index in date_columns:
            for cells in page.iter_rows(min_row=5, min_col=index, max_col=index):
                cells[0].number_format = "mm/dd/yy"
        for index in decimal_columns:
            for cells in page.iter_rows(min_row=5, min_col=index, max_col=index):
                cells[0].number_format = "0.00"
        if values:
            table = Table(displayName=name.replace(" ", ""), ref=f"A4:{get_column_letter(len(titles))}{page.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            page.add_table(table)
        else:
            page["A5"] = "No records"
        page.freeze_panes = "C5" if len(titles) > 3 else "A5"
        page.sheet_properties.pageSetUpPr.fitToPage = True
        page.page_setup.orientation = "landscape"
        page.page_setup.fitToWidth = 1
        page.page_setup.fitToHeight = 0
        page.print_title_rows = "1:4"
        return page

    summary = sheet("Summary", ["Measure", "Value"], [
        ["Scope", "COMPLETE WEEK ARCHIVE" if complete else "SCOPED REPORT ONLY"],
        ["Generated UTC", generated_at.isoformat() + "Z"], ["Version", version],
        ["Employees", len(weekly)], ["Sort slices", len(rows)],
        ["Worked hours", float(timecards.hours(sum(item["seconds"] or 0 for item in rows)))],
        ["Exceptions", len(exceptions)],
        ["Previously purged slices", previously_purged_slices],
        ["Time basis", "Exact UTC timestamps; unpaid gaps excluded"],
        ["Attendance", "Only Here contributes worked hours"],
    ])
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 76
    sheet("All Employee Hours", headers, detail, (3,), (10, 11))
    sheet("Employee Weekly", ["Employee ID", "Employee", "Hours", "Attendance statuses"],
        [[*employee_labels[key], float(timecards.hours(value[0])) if "here" in value[1] else None, ", ".join(sorted(value[1]))] for key, value in sorted(weekly.items(), key=lambda entry: employee_labels[entry[0]])], decimal_columns=(3,))
    sheet("Sort Labor", ["Gateway", "Sort", "Operation", "Department", "Work Area", "Hours", "Attendance statuses"],
        [[*key, float(timecards.hours(value[0])) if "here" in value[1] else None, ", ".join(sorted(value[1]))] for key, value in sorted(labor.items())], decimal_columns=(6,))
    sheet("Exceptions", headers[:8] + ["Exception", "Slice ID"], exceptions, (3,))
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    raw_headers = headers[:8] + ["Attendance", "Slice ID", "Source Operation ID", "Version", "Segment ID", "Start UTC", "End UTC",
                                  "Local timezone", "Updated by user ID", "Updated UTC", "Person ID"]
    files = {"timecards.xlsx": xlsx.getvalue(), "timecards.csv": csv_bytes(raw_headers, raw),
             "edit-history.csv": csv_bytes(["Edit ID", "Slice ID", "Version", "Actor ID", "Recorded UTC", "Snapshot JSON"],
                 [[edit.id, edit.slice_id, edit.version, edit.actor_id, edit.recorded_at.isoformat() + "Z", edit.snapshot_json] for edit in edits])}
    manifest = {"format_version": 1, "complete": complete, "week_start": week.isoformat(),
                "week_end": (week + timedelta(days=6)).isoformat(), "generated_utc": generated_at.isoformat() + "Z",
                "data_version": version, "employee_count": len(weekly), "slice_count": len(rows),
                "segment_count": sum(len(item["segments"]) for item in rows), "edit_count": len(edits),
                "previously_purged_slices": previously_purged_slices,
                "sha256": {name: hashlib.sha256(data).hexdigest() for name, data in files.items()}}
    files["manifest.json"] = json.dumps(manifest, sort_keys=True, indent=2).encode()
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return package.getvalue()


def generate(user, week, *, complete, as_of):
    if week != timecards.week_start(week) or week + timedelta(days=6) >= as_of:
        raise ValueError("Choose a completed Sunday–Saturday week.")
    if complete and not can_archive(user):
        raise ValueError("Complete Week Archive requires NeoStaffing Master access.")
    # This same lock is taken by every report-affecting writer. Build a coherent
    # snapshot; generation alone cannot schedule deletion or assert delivery.
    from app.models.staffing_timecard import StaffingTimecardWeek
    receipt = (timecards.lock_week(week) if complete else StaffingTimecardWeek.query.filter_by(
        week_start=week).with_for_update(read=True).one_or_none())
    if receipt and receipt.purged_at:
        raise ValueError("This week has already been purged. Use the downloaded archive.")
    rows = timecards.read_rows(user, week, week + timedelta(days=6), complete=complete)
    edits = Edit.query.filter(Edit.slice_id.in_([row["slice"].id for row in rows])).order_by(Edit.id).all()
    version = receipt.version if receipt else None
    package = build_package(rows, week, version, complete=complete, edits=edits,
                            previously_purged_slices=receipt.purged_slices if receipt else 0)
    token = serializer().dumps({"week": week.isoformat(), "version": version, "actor": user.id,
                              "sha256": hashlib.sha256(package).hexdigest()}) if complete else None
    return package, token


def acknowledge_download(user, token, *, now=None):
    if not can_archive(user):
        raise ValueError("Complete Week Archive requires NeoStaffing Master access.")
    from datetime import date
    try:
        payload = serializer().loads(token, max_age=3600)
    except BadSignature as error:
        raise ValueError("Archive acknowledgement expired. Download again.") from error
    receipt = timecards.lock_week(date.fromisoformat(payload["week"]))
    if payload["actor"] != user.id or receipt.version != payload["version"] or receipt.purged_at:
        raise ValueError("Archive is stale. Download a new complete archive.")
    if receipt.archive_version == payload["version"] and receipt.archive_sha256 == payload["sha256"] and receipt.purge_after:
        return receipt  # Idempotent acknowledgement; no grace extension on replay.
    now = now or datetime.utcnow()
    receipt.archive_version = payload["version"]
    receipt.archive_sha256 = payload["sha256"]
    receipt.archive_actor_id = user.id
    receipt.downloaded_at = now
    receipt.purge_after = now + timedelta(days=3)
    return receipt
